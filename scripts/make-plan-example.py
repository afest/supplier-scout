"""Пример плана закупа для помощника (T-546): суши-бар на доставке, неделя. Как пишут в реальном закупочном файле —
название с брендом и фасовкой, единица, количество; сверху название документа, внизу итог. Две позиции — вне набора
(овощи), чтобы итог показал, что не закрыто.

Запуск: python scripts/make-plan-example.py   → web/public/plan-sushi-bar.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / "web" / "public" / "plan-sushi-bar.xlsx"

ROWS = [
    ("Рис для суши «Японика» Sen Soy, мешок 25 кг", "мешок", 2, ""),
    ("Нори Gold 50 листов, 140 г", "уп", 10, ""),
    ("Соевый соус Kikkoman 1 л", "бут", 6, ""),
    ("Васаби порошок «Премиум» 1 кг", "уп", 2, ""),
    ("Имбирь маринованный розовый 1,4 кг", "ведро", 4, ""),
    ("Уксус рисовый Mizkan 20 л", "канистра", 1, "раз в месяц"),
    ("Кунжут белый жареный 1 кг", "уп", 2, ""),
    ("Сухари панко 1 кг", "уп", 3, ""),
    ("Лосось филе охл. трим D 1,5–2 кг", "кг", 15, "только охлаждённый"),
    ("Тунец филе (лоин) зам. 1 кг", "кг", 3, ""),
    ("Креветка тигровая 16/20 б/г 1 кг", "уп", 4, ""),
    ("Икра масаго оранжевая 0,5 кг", "уп", 2, ""),
    ("Угорь жареный унаги 0,25 кг", "шт", 6, ""),
    ("Сыр сливочный Hochland 2,2 кг", "шт", 3, ""),
    ("Сливки 33% 1 л", "шт", 4, "для соусов"),
    ("Контейнер для роллов с крышкой, 1 секция", "шт", 1000, ""),
    ("Палочки для суши в инд. упаковке", "шт", 2000, ""),
    ("Пакет крафт с плоскими ручками 32×20", "шт", 500, ""),
    ("Огурцы свежие", "кг", 10, "на рынке"),
    ("Авокадо Хасс", "кг", 5, ""),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "План закупа"
    ws["A1"] = "План закупа на неделю · суши-бар на доставке, Екатеринбург"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Составлен 15.10.2026"
    ws["A2"].font = Font(italic=True, color="808080")
    head = ["№", "Наименование", "Ед.", "Кол-во", "Комментарий"]
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EDEDED")
    for i, (name, unit, qty, note) in enumerate(ROWS, 1):
        ws.append([i, name, unit, qty, note or None])
    ws.append([])
    ws.append([None, f"Итого позиций: {len(ROWS)}"])
    ws[ws.max_row][1].font = Font(bold=True)
    for col, width in zip("ABCDE", (5, 46, 10, 9, 22)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5, max_row=4 + len(ROWS)):
        row[3].alignment = Alignment(horizontal="right")
    wb.properties.creator = "supplier-scout"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
