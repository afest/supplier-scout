"""Загрузка страниц сайта поставщика: robots.txt, текст из HTML/PDF/XLSX, запасной путь через Jina Reader."""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import threading
import time
from datetime import date
from urllib.parse import urljoin, urlparse, urldefrag

import httpx
from lxml import etree, html as lxml_html

from .config import CACHE, CATALOG_TEXT_CAP, JINA_MIN_TEXT, MAX_FETCHES_PER_SITE, PAGE_TEXT_CAP, USER_AGENT

PAGES_DIR = CACHE / "pages"
logging.getLogger("pypdf").setLevel(logging.ERROR)     # предупреждения о шрифтах PDF на текст не влияют, а лог забивают
BOT_TOKEN = "SupplierScoutBot"
BLOCK_TAGS = {"p", "div", "br", "li", "ul", "ol", "tr", "td", "th", "table", "section", "article", "header", "footer",
              "h1", "h2", "h3", "h4", "h5", "h6", "nav", "aside", "main", "form", "dl", "dt", "dd", "address", "blockquote"}
DROP_TAGS = ["script", "style", "noscript", "svg", "template", "iframe", "canvas", "head"]
SKIP_EXT = re.compile(r"\.(jpe?g|png|gif|webp|svg|ico|zip|rar|7z|docx?|pptx?|mp4|mp3|avi|exe|apk)(\?|$)", re.I)


class Robots:
    """robots.txt по RFC 9309: группа своего бота или «*», самое длинное совпадение, Allow при равной длине, `*` и `$`.

    Стандартный urllib.robotparser здесь не годится: `Disallow: /?` («не индексировать адреса с параметрами»)
    он превращает в `Disallow: /` и запрещает весь сайт.
    """

    def __init__(self, text: str):
        self.groups: list[tuple[list[str], list[tuple[bool, str]], float | None]] = []
        self.crawl_delay: float | None = None
        self.block_all = False
        agents, rules, delay, in_rules = [], [], None, False
        for raw in text.splitlines():
            line = raw.split("#", 1)[0].strip()
            if ":" not in line:
                continue
            key, val = (x.strip() for x in line.split(":", 1))
            key = key.lower()
            if key == "user-agent":
                if in_rules:
                    self.groups.append((agents, rules, delay))
                    agents, rules, delay, in_rules = [], [], None, False
                agents.append(val.lower())
            elif key in ("allow", "disallow"):
                in_rules = True
                if val:
                    rules.append((key == "allow", val))
            elif key == "crawl-delay":
                in_rules = True
                try:
                    delay = float(val)
                except ValueError:
                    pass
        if agents:
            self.groups.append((agents, rules, delay))

    @classmethod
    def disallow_all(cls, why: str) -> "Robots":
        rb = cls("")
        rb.block_all = True
        rb.why = why
        return rb

    @staticmethod
    def _match(pattern: str, path: str) -> bool:
        rx = re.escape(pattern).replace(r"\*", ".*")
        if rx.endswith(r"\$"):
            rx = rx[:-2] + "$"
        return re.match(rx, path) is not None

    def allowed(self, token: str, url: str) -> bool:
        if self.block_all:
            return False
        p = urlparse(url)
        path = (p.path or "/") + (f"?{p.query}" if p.query else "")
        token = token.lower()
        mine = [g for g in self.groups if any(a != "*" and a in token for a in g[0])]
        groups = mine or [g for g in self.groups if "*" in g[0]]
        best: tuple[int, bool] | None = None
        for _, rules, delay in groups:
            if delay is not None:
                self.crawl_delay = delay
            for allow, pat in rules:
                if self._match(pat, path) and (best is None or len(pat) > best[0] or (len(pat) == best[0] and allow)):
                    best = (len(pat), allow)
        return best is None or best[1]


def host_of(url: str) -> str:
    h = (urlparse(url).hostname or "").lower()
    return h[4:] if h.startswith("www.") else h


# Однобайтные кодировки декодируют любые байты без ошибки. Сервер часто отдаёт ISO-8859-1 по умолчанию, а тело
# в cp1251 — такой заголовок пробуем последним, иначе кириллица молча станет мусором.
LATIN = {"iso-8859-1", "latin-1", "latin1", "iso8859-1", "us-ascii", "ascii", "windows-1252", "cp1252"}


def _decode(content: bytes, content_type: str) -> str:
    m = re.search(r"charset=([\w-]+)", content_type or "", re.I)
    header = [m.group(1)] if m else []
    head = content[:4000].decode("ascii", "ignore")
    m2 = re.search(r"<meta[^>]+charset=[\"']?([\w-]+)", head, re.I)
    meta = [m2.group(1)] if m2 else []
    if header and header[0].lower() in LATIN:
        candidates = [e for e in meta if e.lower() not in LATIN] + ["utf-8", "cp1251"] + header
    else:
        candidates = header + meta + ["utf-8", "cp1251"]
    for enc in candidates:
        try:
            return content.decode(enc)
        except (LookupError, UnicodeDecodeError):
            continue
    return content.decode("utf-8", "replace")


def _norm_ws(text: str) -> str:
    text = text.replace("\xa0", " ").replace(" ", " ").replace("​", "")
    lines = [re.sub(r"[ \t\r\f\v]+", " ", l).strip() for l in text.split("\n")]
    out, prev = [], None
    for l in lines:
        if not l or l == prev:          # пустые и повторяющиеся подряд строки (меню, дубли блоков)
            continue
        out.append(l)
        prev = l
    return "\n".join(out)


def html_to_text_and_links(raw: str, base_url: str) -> tuple[str, str, list[dict]]:
    doc = lxml_html.fromstring(raw)
    title = (doc.findtext(".//title") or "").strip()
    links, contacts = [], []
    for a in doc.iter("a"):
        href = (a.get("href") or "").strip()
        if not href:
            continue
        label = " ".join((a.text_content() or "").split())[:120]
        low = href.lower()
        if low.startswith(("tel:", "mailto:")) or any(k in low for k in ("wa.me", "whatsapp", "t.me/", "telegram", "viber", "vk.com/", "max.ru/")):
            contacts.append(href)
            continue
        if low.startswith(("javascript:", "#")):
            continue
        links.append({"url": urldefrag(urljoin(base_url, href))[0], "text": label})
    for el in list(doc.iter(*DROP_TAGS)):
        if el.getparent() is not None:
            el.drop_tree()
    root = doc.find(".//body")
    root = root if root is not None else doc
    parts: list[str] = []
    for event, el in etree.iterwalk(root, events=("start", "end")):
        tag = el.tag if isinstance(el.tag, str) else None      # комментарии и инструкции — без текста, только хвост
        if event == "start":
            if tag in BLOCK_TAGS:
                parts.append("\n")
            if tag and el.text:
                parts.append(el.text)
        else:
            if tag in BLOCK_TAGS:
                parts.append("\n")
            if el.tail and el is not root:
                parts.append(el.tail)
    text = _norm_ws("".join(parts))
    if contacts:
        uniq = list(dict.fromkeys(contacts))
        text += "\n[ссылки-контакты на странице] " + " ; ".join(uniq[:30])
    return title, text, links


def pdf_to_text(content: bytes, max_pages: int = 6) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    return _norm_ws("\n".join((p.extract_text() or "") for p in reader.pages[:max_pages]))


def xlsx_rows(content: bytes, max_rows: int = 3000, max_sheets: int = 5) -> list[list]:
    """Строки прайса как есть — для разбора таблицы по колонкам (catalog.xlsx_items); хранятся в кэше страницы."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:max_sheets]:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            out.append([c if isinstance(c, (int, float, str)) or c is None else str(c) for c in row])
        out.append(None)                # граница листа: заголовок таблицы ищется заново
    return out


def xlsx_to_text(content: bytes, max_rows: int = 300, max_sheets: int = 2) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets[:max_sheets]:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                rows.append(" | ".join(cells))
    return _norm_ws("\n".join(rows))


ANTIBOT = (401, 403, 429)       # сайт отказал боту — обходить через посредника не будем


class Fetcher:
    """Вежливая загрузка: robots.txt, пауза между запросами к одному сайту, общий лимит загрузок на сайт, кэш на диске.

    offline=True — только кэш: в сеть не ходит, страницы, которых нет в кэше, возвращаются с ошибкой.
    """

    def __init__(self, delay: float = 1.0, offline: bool = False, max_per_site: int = MAX_FETCHES_PER_SITE):
        self.http = httpx.Client(follow_redirects=True, timeout=httpx.Timeout(45, connect=12),
                                 headers={"User-Agent": USER_AGENT, "Accept-Language": "ru-RU,ru;q=0.9"},
                                 verify=True, trust_env=False)
        self.delay = delay
        self.offline = offline
        self.max_per_site = max_per_site
        self.fetched: dict[str, int] = {}         # загрузки из сети за этот прогон по сайтам — против лимита
        self.robots: dict[str, Robots] = {}
        self.last_hit: dict[str, float] = {}
        self.lock = threading.Lock()

    # ---- robots.txt ----
    def _robots(self, url: str) -> "Robots":
        p = urlparse(url)
        key = f"{p.scheme}://{p.netloc}"
        if key not in self.robots:
            try:
                r = self.http.get(f"{key}/robots.txt", timeout=12)
                if r.status_code >= 500:
                    rb = Robots.disallow_all("robots.txt отвечает ошибкой сервера")
                elif r.status_code == 200 and "text/html" not in r.headers.get("content-type", ""):
                    rb = Robots(_decode(r.content, r.headers.get("content-type", "")))
                else:
                    rb = Robots("")         # 4xx или HTML вместо файла — правил нет (RFC 9309, 2.3.1.3)
            except httpx.HTTPError:
                # RFC 9309, 2.3.1.4: robots.txt недоступен из-за сети или сервера — сайт считается закрытым целиком
                rb = Robots.disallow_all("robots.txt не открылся")
            self.robots[key] = rb
        return self.robots[key]

    def allowed(self, url: str) -> bool:
        return self._robots(url).allowed(BOT_TOKEN, url)

    def _wait(self, host: str):
        rb = next((v for k, v in self.robots.items() if urlparse(k).hostname and host_of(k) == host), None)
        delay = max(self.delay, (rb.crawl_delay if rb else 0) or 0)
        delay = min(delay, 10.0)
        with self.lock:
            last = self.last_hit.get(host, 0)
            wait = last + delay - time.time()
            self.last_hit[host] = max(time.time(), last + delay)
        if wait > 0:
            time.sleep(wait)

    # ---- одна страница ----
    def page(self, url: str, use_cache: bool = True, full: bool = False) -> dict:
        """full=True — для каталога и прайс-листов: больше текста, до 40 страниц PDF, до 3000 строк таблицы."""
        key = hashlib.sha1(url.encode()).hexdigest()[:16]
        path = PAGES_DIR / host_of(url) / f"{key}.json"
        attempts = 0
        if use_cache and path.exists():
            cached = self._usable(json.loads(path.read_text(encoding="utf-8")))
            attempts = int(cached.get("attempts") or (2 if cached.get("retried") else 1))
            short = full and not cached.get("full") and (cached.get("kind") in ("pdf", "xlsx") or len(cached.get("text", "")) >= PAGE_TEXT_CAP * 2)
            if self.offline or ((not cached.get("error") or attempts >= 3) and not short):   # неудачную загрузку повторяем, всего до трёх попыток
                return cached
        res = {"url": url, "final_url": url, "status": None, "fetched_at": date.today().isoformat(),
               "title": "", "text": "", "links": [], "via": "direct", "error": None, "kind": "html", "attempts": attempts + 1,
               "full": full}
        if self.offline:
            res["error"] = "нет в кэше, а прогон офлайн"
            return res
        if not self.allowed(url):
            rb = self._robots(url)
            res["error"] = f"{rb.why} — сайт закрыт для сбора (RFC 9309)" if rb.block_all else "robots.txt запрещает"
            return self._save(path, res)
        host = host_of(url)
        with self.lock:
            if self.fetched.get(host, 0) >= self.max_per_site:
                res["error"] = f"лимит {self.max_per_site} загрузок на сайт за прогон"
                return res
            self.fetched[host] = self.fetched.get(host, 0) + 1
        self._wait(host)
        try:
            r = self.http.get(url)
            res["status"], res["final_url"] = r.status_code, str(r.url)
            ctype = r.headers.get("content-type", "").lower()
            if r.status_code >= 400:
                res["error"] = f"HTTP {r.status_code}"
            elif "pdf" in ctype or url.lower().endswith(".pdf"):
                res["kind"], res["text"] = "pdf", pdf_to_text(r.content, max_pages=40 if full else 6)
            elif "spreadsheet" in ctype or "excel" in ctype or re.search(r"\.xlsx?($|\?)", url, re.I):
                res["kind"], res["text"] = "xlsx", xlsx_to_text(r.content, max_rows=3000 if full else 300, max_sheets=5 if full else 2)
                if full:
                    res["rows"] = xlsx_rows(r.content)
            else:
                res["title"], res["text"], res["links"] = html_to_text_and_links(_decode(r.content, ctype), str(r.url))
        except Exception as e:  # сеть, сертификат, битый HTML — фиксируем причину, сбор продолжается
            res["error"] = f"{type(e).__name__}: {str(e)[:160]}"
        # Jina — только для сайтов на JS: страница отдалась, а текста почти нет. После отказа сервера (403, 429…)
        # или сетевой ошибки посредника не зовём — это был бы обход запрета.
        if len(res["text"]) < JINA_MIN_TEXT and res["kind"] == "html" and not res["error"] and res["status"] and res["status"] < 400:
            self._jina(url, res)
        res["text"] = res["text"][: (CATALOG_TEXT_CAP if full else PAGE_TEXT_CAP * 2)]
        return self._save(path, res)

    @staticmethod
    def _usable(p: dict) -> dict:
        """Страница из кэша, полученная через Jina после отказа сервера (так собирали до 14.10), — не используется."""
        if p.get("via") == "jina" and p.get("status") in ANTIBOT:
            return {**p, "text": "", "links": [], "error": f"HTTP {p['status']}: сайт отказал, посредник не используется", "attempts": 3}
        return p

    def _jina(self, url: str, res: dict):
        try:
            r = httpx.get(f"https://r.jina.ai/{url}", timeout=60, headers={"User-Agent": USER_AGENT, "X-Return-Format": "text"})
            if r.status_code == 200 and len(r.text) > len(res["text"]):
                res["text"], res["via"], res["error"] = _norm_ws(r.text), "jina", None
        except httpx.HTTPError as e:
            res["error"] = f"Jina: {type(e).__name__}"

    @staticmethod
    def cached(url: str) -> dict | None:
        """Страница из кэша по исходному или конечному адресу (после редиректа)."""
        folder = PAGES_DIR / host_of(url)
        for f in folder.glob("*.json") if folder.exists() else []:
            p = Fetcher._usable(json.loads(f.read_text(encoding="utf-8")))
            if url in (p.get("url"), p.get("final_url")) and p.get("text"):
                return p
        return None

    @staticmethod
    def _save(path, res: dict) -> dict:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
        return res


def pick_pages(home: dict, extra_urls: list[str], hints: list[tuple[str, int]], limit: int) -> list[str]:
    """Внутренние страницы по словам в ссылке и её тексте: контакты, доставка, реквизиты, опт, прайс…"""
    base_host = host_of(home["final_url"] or home["url"])
    scored: dict[str, int] = {}
    for link in home.get("links", []):
        u = link["url"]
        if host_of(u) != base_host or SKIP_EXT.search(u) or u.rstrip("/") == (home["final_url"] or "").rstrip("/"):
            continue
        hay = (link["text"] + " " + u).lower()
        score = sum(w for k, w in hints if k in hay)
        if score:
            scored[u] = max(scored.get(u, 0), score)
    picked = [u for u in extra_urls if host_of(u) == base_host][:1]   # страница из выдачи поиска — про нужный товар
    seen = {u.rstrip("/") for u in picked} | {(home["final_url"] or "").rstrip("/")}
    for u, _ in sorted(scored.items(), key=lambda kv: -kv[1]):
        if len(picked) >= limit:
            break
        if u.rstrip("/") not in seen:
            picked.append(u)
            seen.add(u.rstrip("/"))
    return picked[:limit]
